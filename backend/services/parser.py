import pdfplumber
import json
import re
from pathlib import Path
from docx import Document
from typing import Dict, Any, Optional


def parse_pdf(file_path: str) -> Dict[str, Any]:
    """Extract text and structured data from PDF."""
    text = ""
    pages = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text() or ""
                pages.append({"page": i + 1, "text": page_text})
                text += page_text + "\n"
    except Exception as e:
        raise ValueError(f"Failed to parse PDF: {str(e)}")
    return {"raw_text": text.strip(), "pages": pages, "format": "pdf"}


def parse_docx(file_path: str) -> Dict[str, Any]:
    """Extract text and structured data from DOCX."""
    text = ""
    paragraphs = []
    tables = []
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())
                text += para.text + "\n"
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            tables.append(table_data)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {str(e)}")
    return {"raw_text": text.strip(), "paragraphs": paragraphs, "tables": tables, "format": "docx"}


def extract_trade_fields(text: str) -> Dict[str, Any]:
    """Extract structured trade fields from raw text using regex patterns."""
    fields = {}

    # Trade type detection
    trade_type_patterns = {
        "fx_spot": r"\b(fx\s*spot|foreign\s*exchange\s*spot|spot\s*fx|spot\s*transaction)\b",
        "fx_forward": r"\b(fx\s*forward|forward\s*fx|non.deliverable\s*forward|ndf|forward\s*contract)\b",
        "interest_rate_swap": r"\b(interest\s*rate\s*swap|irs|rate\s*swap|fixed.floating\s*swap)\b",
        "equity_trade": r"\b(equity\s*trade|share\s*purchase|stock\s*trade|equity\s*transaction)\b",
        "bond_trade": r"\b(bond\s*trade|fixed\s*income|debt\s*security|bond\s*purchase|note\s*purchase)\b",
        "credit_default_swap": r"\b(credit\s*default\s*swap|cds|credit\s*derivative|protection\s*buyer)\b",
    }
    for trade_type, pattern in trade_type_patterns.items():
        if re.search(pattern, text, re.IGNORECASE):
            fields["trade_type"] = trade_type
            break

    # Numeric fields
    patterns = {
        "notional_amount": r"notional\s*(?:amount|principal)?[:\s]*(?:USD|EUR|GBP|JPY|SGD|INR)?\s*([\d,]+(?:\.\d+)?)",
        "exchange_rate": r"(?:exchange\s*rate|spot\s*rate|rate)[:\s]*([\d.]+)",
        "forward_rate": r"forward\s*rate[:\s]*([\d.]+)",
        "fixed_rate": r"fixed\s*rate[:\s]*([\d.]+)%?",
        "premium_rate": r"(?:premium|spread|running\s*spread)[:\s]*([\d.]+)%?",
        "price_per_share": r"(?:price|share\s*price)[:\s]*(?:USD|EUR|GBP)?\s*([\d.]+)",
        "number_of_shares": r"(?:shares?|quantity)[:\s]*([\d,]+)",
        "coupon_rate": r"coupon\s*rate[:\s]*([\d.]+)%?",
    }
    for field, pattern in patterns.items():
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            val = match.group(1).replace(",", "")
            try:
                fields[field] = float(val)
            except:
                fields[field] = val

    # Date fields
    date_patterns = {
        "trade_date": r"trade\s*date[:\s]*([\d]{4}-[\d]{2}-[\d]{2}|[\d]{1,2}[/-][\d]{1,2}[/-][\d]{2,4})",
        "value_date": r"(?:value|settlement)\s*date[:\s]*([\d]{4}-[\d]{2}-[\d]{2}|[\d]{1,2}[/-][\d]{1,2}[/-][\d]{2,4})",
        "maturity_date": r"(?:maturity|termination|expiry)\s*date[:\s]*([\d]{4}-[\d]{2}-[\d]{2}|[\d]{1,2}[/-][\d]{1,2}[/-][\d]{2,4})",
    }
    for field, pattern in date_patterns.items():
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            fields[field] = match.group(1)

    # String fields
    string_patterns = {
        "counterparty": r"counterparty[:\s]+([A-Z][A-Za-z\s&.,]+(?:Bank|AG|PLC|LLC|Ltd|NA|Inc\.?)?)",
        "currency_pair": r"currency\s*pair[:\s]*([A-Z]{3}[/][A-Z]{3})",
        "reference_entity": r"reference\s*entity[:\s]+([A-Za-z\s]+(?:Inc\.?|Corp\.?|Ltd\.?|PLC)?)",
        "tenor": r"tenor[:\s]*([\d]+[YMW])",
        "isin": r"ISIN[:\s]*([A-Z]{2}[A-Z0-9]{10})",
    }
    for field, pattern in string_patterns.items():
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            fields[field] = match.group(1).strip()

    return fields


def parse_document(file_path: str) -> Dict[str, Any]:
    """Main entry point: parse any supported document format."""
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        parsed = parse_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        parsed = parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")
    parsed["extracted_fields"] = extract_trade_fields(parsed["raw_text"])
    return parsed